#!/usr/bin/env python3
"""
build_excel_mirror.py — Genera el Excel espejo (human-readable) de un JSON de
extracción de categoría (esquema v3, ver references/schema.md), resaltando
visualmente las celdas cuyo campo original tiene `confianza: baja`.

El Excel se genera SIEMPRE a partir del JSON, nunca al revés (regla de
references/schema.md, sección "Formato de salida y versionado"): este script
solo lee y formatea, no reinterpreta ni corrige ningún valor.

Uso:
    python3 build_excel_mirror.py --json Extracciones/categoria_contratos_v3.json
    python3 build_excel_mirror.py --json ruta/v3.json --output ruta/custom.xlsx

Requiere openpyxl (ya presente en el entorno de este skill).
"""
import argparse
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FILL_BAJA = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
FONT_BAJA = Font(color="9C6500")
FONT_HEADER = Font(bold=True, color="FFFFFF")
FILL_HEADER = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")

# (encabezado, ruta de acceso dentro del contrato — tupla de claves para navegar dicts anidados)
COLUMNS = [
    ("numero_contrato", ("numero_contrato",)),
    ("proveedor", ("proveedor",)),
    ("rut_run_proveedor", ("rut_run_proveedor",)),
    ("vigencia.fecha_inicio", ("vigencia", "fecha_inicio")),
    ("vigencia.fecha_fin", ("vigencia", "fecha_fin")),
    ("vigencia.renovacion_automatica_tacita", ("vigencia", "renovacion_automatica_tacita")),
    ("vigencia.termino_anticipado", ("vigencia", "termino_anticipado")),
    ("monto_total", ("monto_total",)),
    ("tipo_cambio", ("tipo_cambio",)),
    ("productos_servicios_costo_unitario", ("productos_servicios_costo_unitario",)),
    ("metodo_pago", ("metodo_pago",)),
    ("niveles_servicio", ("niveles_servicio",)),
    ("clausula_incentivo_desempeno", ("clausula_incentivo_desempeno",)),
    ("clausula_multa_penalizacion", ("clausula_multa_penalizacion",)),
    ("mecanismo_reajuste_indexacion", ("mecanismo_reajuste_indexacion",)),
    ("documentos_fuente", ("documentos_fuente",)),
    ("notas", ("notas",)),
]

WIDE_COLUMNS = {
    "notas", "documentos_fuente", "productos_servicios_costo_unitario",
    "monto_total", "clausula_multa_penalizacion",
}


def _get_by_path(contrato: dict, path: tuple):
    value = contrato
    for key in path:
        if not isinstance(value, dict):
            return None
        value = value.get(key)
    return value


def _stringify(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Sí" if value else "No"
    if isinstance(value, (int, float)):
        return f"{value:,}" if isinstance(value, (int, float)) and not isinstance(value, bool) else str(value)
    if isinstance(value, str):
        return value
    if isinstance(value, list):
        return "; ".join(_stringify(v) for v in value if v not in (None, ""))
    if isinstance(value, dict):
        if {"moneda", "valor", "periodo"} <= value.keys():
            monto = value.get("valor")
            monto_str = f"{monto:,}" if isinstance(monto, (int, float)) else str(monto)
            return f"{value.get('moneda', '')} {monto_str} ({value.get('periodo', '')})"
        if "descripcion" in value and "cantidad" in value:
            return f"{value.get('descripcion')} (cant: {value.get('cantidad')})"
        if "aplica_a" in value and "valor" in value:
            return f"{value.get('aplica_a')}: {_stringify(value.get('valor'))}"
        if "tipo" in value and "gatillante" in value:
            partes = [value.get("tipo"), value.get("gatillante")]
            if value.get("monto"):
                partes.append(_stringify(value.get("monto")))
            return " / ".join(p for p in partes if p)
        if "valor" in value:
            detalle = value.get("detalle")
            base = _stringify(value.get("valor"))
            if detalle not in (None, ""):
                return f"{base} — {_stringify(detalle)}"
            return base
        # dict "compuesto" sin 'valor' propio (ej. productos_servicios_costo_unitario)
        partes = []
        for k, v in value.items():
            if k in ("fuente", "confianza"):
                continue
            texto = _stringify(v)
            if texto:
                partes.append(f"{k}: {texto}" if not isinstance(v, list) else texto)
        return "; ".join(partes)
    return str(value)


def _confianza_of(value) -> str | None:
    if isinstance(value, dict):
        if "confianza" in value:
            return value.get("confianza")
        peor = None
        orden = {"baja": 3, "media": 2, "alta": 1}
        for v in value.values():
            c = _confianza_of(v)
            if c and (peor is None or orden.get(c, 0) > orden.get(peor, 0)):
                peor = c
        return peor
    if isinstance(value, list):
        peor = None
        orden = {"baja": 3, "media": 2, "alta": 1}
        for v in value:
            c = _confianza_of(v)
            if c and (peor is None or orden.get(c, 0) > orden.get(peor, 0)):
                peor = c
        return peor
    return None


def build_workbook(data: dict) -> Workbook:
    wb = Workbook()

    # --- Hoja Metadata ---
    ws_meta = wb.active
    ws_meta.title = "Metadata"
    metadata = data.get("metadata", {})
    ws_meta.append(["Campo", "Valor"])
    for cell in ws_meta[1]:
        cell.font = FONT_HEADER
        cell.fill = FILL_HEADER
    for key in ("categoria", "fecha_extraccion", "cliente", "metodologia", "version",
                "version_anterior", "total_contratos", "nota_normalizacion"):
        ws_meta.append([key, _stringify(metadata.get(key))])
    ws_meta.column_dimensions["A"].width = 22
    ws_meta.column_dimensions["B"].width = 90

    # --- Hoja Contratos ---
    ws = wb.create_sheet("Contratos")
    ws.append(["Leyenda: celdas resaltadas = confianza: baja (revisar contra el documento fuente)"])
    ws["A1"].font = Font(italic=True, color="9C6500")
    ws.append([col_name for col_name, _ in COLUMNS])
    header_row = 2
    for cell in ws[header_row]:
        cell.font = FONT_HEADER
        cell.fill = FILL_HEADER
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    contratos = data.get("contratos", [])
    for contrato in contratos:
        row_values = []
        row_confianzas = []
        for _, path in COLUMNS:
            raw = _get_by_path(contrato, path)
            row_values.append(_stringify(raw))
            row_confianzas.append(_confianza_of(raw))
        ws.append(row_values)
        row_idx = ws.max_row
        for col_idx, confianza in enumerate(row_confianzas, start=1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if confianza == "baja":
                cell.fill = FILL_BAJA
                cell.font = FONT_BAJA

    for col_idx, (col_name, _) in enumerate(COLUMNS, start=1):
        letter = get_column_letter(col_idx)
        ws.column_dimensions[letter].width = 45 if col_name in WIDE_COLUMNS else 22

    ws.freeze_panes = f"A{header_row + 1}"

    return wb


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--json", required=True, help="Ruta al JSON de extracción de categoría (v3)")
    parser.add_argument("--output", default=None, help="Ruta de salida .xlsx (default: mismo nombre que el JSON, extensión .xlsx)")
    args = parser.parse_args()

    json_path = Path(args.json)
    if not json_path.exists():
        print(f"❌ No existe el archivo: {json_path}", file=sys.stderr)
        sys.exit(1)

    try:
        data = json.loads(json_path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as e:
        print(f"❌ JSON inválido en {json_path}: {e}", file=sys.stderr)
        sys.exit(1)

    if "contratos" not in data:
        print(f"❌ {json_path} no tiene la forma esperada (falta 'contratos' a nivel raíz)", file=sys.stderr)
        sys.exit(1)

    output_path = Path(args.output) if args.output else json_path.with_suffix(".xlsx")

    wb = build_workbook(data)
    wb.save(output_path)

    n_baja = sum(
        1
        for contrato in data.get("contratos", [])
        for _, path in COLUMNS
        if _confianza_of(_get_by_path(contrato, path)) == "baja"
    )
    print(f"✅ Excel generado: {output_path} ({len(data.get('contratos', []))} contratos, {n_baja} celdas con confianza: baja resaltadas)")


if __name__ == "__main__":
    main()
